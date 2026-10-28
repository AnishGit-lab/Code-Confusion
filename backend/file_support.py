"""
PriviScan File Support
======================

Supports text extraction and privacy analysis for:

    PDF
    DOCX
    PPTX
    XLSX
    TXT
    CSV
    JSON
    MD

Extracted text is passed into the existing PriviScan text-analysis
pipeline.

For PDF/DOCX/PPTX files, embedded or rendered images can also be
analyzed when the required libraries are available.

Risk scoring remains centralized in backend.risk.scoring.
This module does not define individual sensitive-information weights.
"""

from __future__ import annotations

import csv
import json
import math
import mimetypes
import re
import tempfile
import zipfile

from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional


# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

SUPPORTED_EXTENSIONS = {
    ".pdf": "PDF",
    ".docx": "DOCX",
    ".pptx": "PPTX",
    ".xlsx": "XLSX",
    ".txt": "TXT",
    ".csv": "CSV",
    ".json": "JSON",
    ".md": "MARKDOWN",
}


# ============================================================
# OPTIONAL LIBRARIES
# ============================================================

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None


# ============================================================
# BASIC HELPERS
# ============================================================

def normalize_whitespace(text: str) -> str:
    """
    Normalize repeated spaces while preserving line structure.
    """
    if text is None:
        return ""

    lines = []

    for line in str(text).splitlines():
        line = re.sub(r"[ \t]+", " ", line).strip()

        if line:
            lines.append(line)

    return "\n".join(lines)


def get_extension(file_path: str | Path) -> str:
    """Return a lowercase file extension."""
    return Path(file_path).suffix.lower()


def get_file_type(file_path: str | Path) -> str:
    """Return the supported PriviScan file type."""
    extension = get_extension(file_path)

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: "
            f"{extension or '(no extension)'}"
        )

    return SUPPORTED_EXTENSIONS[extension]


def is_supported_file(file_path: str | Path) -> bool:
    """Check whether the file extension is supported."""
    return get_extension(file_path) in SUPPORTED_EXTENSIONS


def get_mime_type(file_path: str | Path) -> str:
    """Return a MIME type where possible."""
    mime_type, _ = mimetypes.guess_type(str(file_path))

    if mime_type:
        return mime_type

    fallback = {
        ".pdf": "application/pdf",
        ".docx": (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        ".pptx": (
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        ),
        ".xlsx": (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        ".txt": "text/plain",
        ".csv": "text/csv",
        ".json": "application/json",
        ".md": "text/markdown",
    }

    return fallback.get(
        get_extension(file_path),
        "application/octet-stream",
    )


def _safe_float(value: Any, default: float = 0.0) -> float:
    """
    Convert a value to a finite float.

    Invalid values, NaN, and infinity use the supplied default.
    """
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError):
        return default

    if not math.isfinite(result):
        return default

    return result


def _clamp_score(value: Any) -> float:
    """Keep a risk score within the supported 0-100 range."""
    return max(0.0, min(100.0, _safe_float(value)))


def _risk_level(score: Any) -> str:
    """
    Derive the risk level from the final risk score.

    Thresholds:
        0-20   LOW
        >20-40 MEDIUM
        >40-70 HIGH
        >70-100 CRITICAL
    """
    score = _clamp_score(score)

    if score <= 20:
        return "LOW"

    if score <= 40:
        return "MEDIUM"

    if score <= 70:
        return "HIGH"

    return "CRITICAL"


def _get_detections(result: Any) -> List[Dict[str, Any]]:
    """
    Extract detections from a dictionary or list result.

    Supports common PriviScan result formats.
    """
    if isinstance(result, list):
        return [
            item
            for item in result
            if isinstance(item, dict)
        ]

    if isinstance(result, dict):
        detections = result.get("detections", [])

        if isinstance(detections, list):
            return [
                item
                for item in detections
                if isinstance(item, dict)
            ]

    return []


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf_text(file_path: str | Path) -> Dict[str, Any]:
    """
    Extract selectable text from a PDF.

    Returns the extracted text and page-level information.
    Scanned PDFs may require OCR through the image-analysis pipeline.
    """
    if PdfReader is None:
        raise RuntimeError(
            "PDF support requires pypdf. "
            "Install it with: pip install pypdf"
        )

    path = Path(file_path)
    reader = PdfReader(str(path))

    pages = []
    all_text = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = ""
            print(
                f"WARNING: Could not extract PDF page "
                f"{page_number}: {exc}"
            )

        text = normalize_whitespace(text)

        pages.append({
            "page": page_number,
            "text": text,
        })

        if text:
            all_text.append(
                f"[PDF Page {page_number}]\n{text}"
            )

    return {
        "text": "\n\n".join(all_text),
        "pages": pages,
        "page_count": len(reader.pages),
    }


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_text(file_path: str | Path) -> Dict[str, Any]:
    """
    Extract text from DOCX paragraphs and tables.
    """
    if Document is None:
        raise RuntimeError(
            "DOCX support requires python-docx. "
            "Install it with: pip install python-docx"
        )

    document = Document(str(file_path))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    table_count = 0

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):
        table_count += 1
        parts.append(f"[DOCX Table {table_index}]")

        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(values)

            if row_text.strip():
                parts.append(row_text)

    return {
        "text": normalize_whitespace("\n".join(parts)),
        "paragraph_count": len(document.paragraphs),
        "table_count": table_count,
    }


# ============================================================
# PPTX EXTRACTION
# ============================================================

def extract_pptx_text(file_path: str | Path) -> Dict[str, Any]:
    """
    Extract text from PPTX slides, shapes, and tables.
    """
    if Presentation is None:
        raise RuntimeError(
            "PPTX support requires python-pptx. "
            "Install it with: pip install python-pptx"
        )

    presentation = Presentation(str(file_path))

    slides = []
    all_text = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_parts = []

        for shape in slide.shapes:
            # Extract ordinary text frames.
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()

                if text:
                    slide_parts.append(text)

            # Extract tables separately.
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    row_text = " | ".join(values)

                    if row_text.strip():
                        slide_parts.append(row_text)

        slide_text = normalize_whitespace(
            "\n".join(slide_parts)
        )

        slides.append({
            "slide": slide_number,
            "text": slide_text,
        })

        if slide_text:
            all_text.append(
                f"[PPTX Slide {slide_number}]\n{slide_text}"
            )

    return {
        "text": "\n\n".join(all_text),
        "slides": slides,
        "slide_count": len(presentation.slides),
    }


# ============================================================
# XLSX EXTRACTION
# ============================================================

def extract_xlsx_text(file_path: str | Path) -> Dict[str, Any]:
    """
    Extract cell values from XLSX worksheets.

    Empty cells are preserved as separators so that adjacent values
    do not accidentally become one continuous string.
    """
    if load_workbook is None:
        raise RuntimeError(
            "XLSX support requires openpyxl. "
            "Install it with: pip install openpyxl"
        )

    workbook = load_workbook(
        filename=str(file_path),
        read_only=True,
        data_only=True,
    )

    sheets = []
    all_text = []

    try:
        for worksheet in workbook.worksheets:
            sheet_parts = []

            for row in worksheet.iter_rows(values_only=True):
                values = []

                for value in row:
                    if value is None:
                        values.append("")
                    else:
                        values.append(str(value).strip())

                # Keep column boundaries for text detection.
                row_text = " | ".join(values).strip(" |")

                if row_text:
                    sheet_parts.append(row_text)

            sheet_text = normalize_whitespace(
                "\n".join(sheet_parts)
            )

            sheets.append({
                "sheet": worksheet.title,
                "text": sheet_text,
            })

            if sheet_text:
                all_text.append(
                    f"[XLSX Sheet: {worksheet.title}]\n"
                    f"{sheet_text}"
                )
    finally:
        workbook.close()

    return {
        "text": "\n\n".join(all_text),
        "sheets": sheets,
        "sheet_count": len(sheets),
    }


# ============================================================
# TXT / MD EXTRACTION
# ============================================================

def extract_plain_text(file_path: str | Path) -> Dict[str, Any]:
    """Extract text using common UTF-8 and international encodings."""
    path = Path(file_path)

    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp949",
        "latin-1",
    ]

    last_error = None

    for encoding in encodings:
        try:
            text = path.read_text(encoding=encoding)

            return {
                "text": normalize_whitespace(text),
                "encoding": encoding,
            }

        except UnicodeDecodeError as exc:
            last_error = exc

    raise RuntimeError(
        f"Could not decode text file {path}: {last_error}"
    )


# ============================================================
# CSV EXTRACTION
# ============================================================

def extract_csv_text(file_path: str | Path) -> Dict[str, Any]:
    """
    Extract CSV rows as readable text.

    CSV headers and values are retained so that the existing
    PriviScan text detector can inspect both field names and values.
    """
    path = Path(file_path)

    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp949",
        "latin-1",
    ]

    last_error = None

    for encoding in encodings:
        try:
            with path.open(
                "r",
                encoding=encoding,
                newline="",
            ) as file:
                reader = csv.reader(file)

                rows = []

                for row in reader:
                    values = [
                        str(value).strip()
                        for value in row
                    ]

                    if not any(values):
                        continue

                    rows.append(" | ".join(values))

            return {
                "text": normalize_whitespace("\n".join(rows)),
                "row_count": len(rows),
                "encoding": encoding,
            }

        except UnicodeDecodeError as exc:
            last_error = exc

    raise RuntimeError(
        f"Could not decode CSV file {path}: {last_error}"
    )


# ============================================================
# JSON EXTRACTION
# ============================================================

def extract_json_text(file_path: str | Path) -> Dict[str, Any]:
    """
    Extract JSON as readable text.

    Valid JSON is formatted before analysis. If the JSON is malformed,
    its raw text is still returned for the text detector.
    """
    path = Path(file_path)

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp949",
        "latin-1",
    ]

    last_error = None

    for encoding in encodings:
        try:
            raw = path.read_text(encoding=encoding)

        except UnicodeDecodeError as exc:
            last_error = exc
            continue

        try:
            data = json.loads(raw)

            # Preserve keys and values in a readable structure.
            formatted = json.dumps(
                data,
                ensure_ascii=False,
                indent=2,
            )

            return {
                "text": normalize_whitespace(formatted),
                "json_valid": True,
                "encoding": encoding,
            }

        except json.JSONDecodeError:
            return {
                "text": normalize_whitespace(raw),
                "json_valid": False,
                "encoding": encoding,
            }

    raise RuntimeError(
        f"Could not decode JSON file {path}: {last_error}"
    )


# ============================================================
# EMBEDDED IMAGE EXTRACTION FROM DOCX/PPTX
# ============================================================

def extract_embedded_images(
    file_path: str | Path,
) -> List[Dict[str, Any]]:
    """
    Extract embedded image files from DOCX/PPTX archives.
    """
    extension = get_extension(file_path)

    if extension not in {".docx", ".pptx"}:
        return []

    image_items = []

    try:
        with zipfile.ZipFile(file_path, "r") as archive:
            for name in archive.namelist():
                if not (
                    name.startswith("word/media/")
                    or name.startswith("ppt/media/")
                ):
                    continue

                if not name.lower().endswith((
                    ".png",
                    ".jpg",
                    ".jpeg",
                    ".webp",
                    ".bmp",
                    ".gif",
                    ".tif",
                    ".tiff",
                )):
                    continue

                try:
                    image_items.append({
                        "name": Path(name).name,
                        "archive_path": name,
                        "bytes": archive.read(name),
                    })

                except Exception as exc:
                    print(
                        f"WARNING: Could not extract embedded "
                        f"image {name}: {exc}"
                    )

    except zipfile.BadZipFile:
        print(
            f"WARNING: {file_path} is not a valid Office ZIP document."
        )

    return image_items


# ============================================================
# PDF PAGE RENDERING
# ============================================================

def render_pdf_pages(
    file_path: str | Path,
    output_directory: str | Path,
    dpi: int = 150,
) -> List[Path]:
    """
    Render PDF pages to PNG files for OCR.

    This is useful when a PDF contains scanned images rather than
    selectable text.
    """
    if fitz is None:
        return []

    if dpi <= 0:
        raise ValueError("PDF rendering DPI must be greater than zero.")

    output_directory = Path(output_directory)
    output_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    output_paths = []
    document = fitz.open(str(file_path))

    try:
        scale = dpi / 72.0
        matrix = fitz.Matrix(scale, scale)

        for page_number in range(len(document)):
            page = document.load_page(page_number)

            pixmap = page.get_pixmap(
                matrix=matrix,
                alpha=False,
            )

            output_path = (
                output_directory / f"page_{page_number + 1}.png"
            )

            pixmap.save(str(output_path))
            output_paths.append(output_path)

    finally:
        document.close()

    return output_paths


# ============================================================
# TEXT EXTRACTION DISPATCHER
# ============================================================

def extract_file_text(
    file_path: str | Path,
) -> Dict[str, Any]:
    """Extract text according to the file type."""
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    if not path.is_file():
        raise ValueError(f"Not a file: {path}")

    file_type = get_file_type(path)

    extractors = {
        "PDF": extract_pdf_text,
        "DOCX": extract_docx_text,
        "PPTX": extract_pptx_text,
        "XLSX": extract_xlsx_text,
        "CSV": extract_csv_text,
        "JSON": extract_json_text,
        "TXT": extract_plain_text,
        "MARKDOWN": extract_plain_text,
    }

    result = extractors[file_type](path)

    result["file_path"] = str(path)
    result["file_name"] = path.name
    result["file_type"] = file_type
    result["extension"] = path.suffix.lower()
    result["mime_type"] = get_mime_type(path)

    return result


# ============================================================
# DETECTION NORMALIZATION
# ============================================================

def normalize_detection(
    detection: Dict[str, Any],
    source: str,
    source_location: Optional[str] = None,
) -> Dict[str, Any]:
    """Normalize a detection from text or image analysis."""
    result = dict(detection)

    result.setdefault("type", "UNKNOWN")
    result.setdefault("value", "")
    result.setdefault("confidence", 0.0)
    result.setdefault("box", None)

    result["source"] = source

    if source_location is not None:
        result["source_location"] = source_location

    result["confidence"] = max(
        0.0,
        min(1.0, _safe_float(result.get("confidence", 0.0))),
    )

    return result


def merge_detections(
    detections: Iterable[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """
    Remove exact duplicate detections while retaining the highest-
    confidence version.

    Source and source location remain part of the key because the same
    value can occur in different pages or embedded images.
    """
    best: Dict[tuple, Dict[str, Any]] = {}

    for detection in detections:
        if not isinstance(detection, dict):
            continue

        dtype = str(detection.get("type", "")).upper()
        value = str(detection.get("value", "")).strip()
        source = str(detection.get("source", ""))
        source_location = str(
            detection.get("source_location", "")
        )

        key = (
            dtype,
            value.casefold(),
            source,
            source_location,
        )

        confidence = _safe_float(
            detection.get("confidence", 0.0)
        )

        if key not in best:
            best[key] = detection
            continue

        old_confidence = _safe_float(
            best[key].get("confidence", 0.0)
        )

        if confidence > old_confidence:
            best[key] = detection

    return list(best.values())


# ============================================================
# CENTRALIZED RISK SCORING
# ============================================================

def _score_detections_compat(
    detections: List[Dict[str, Any]],
) -> Dict[str, Any]:
    """
    Use the centralized risk scorer and normalize its output.

    Supports either:
        1. A dictionary containing detections and aggregate scores.
        2. A list of scored detection dictionaries.

    Risk level is always recalculated from the final risk score.
    A stale risk_level supplied by another function is not trusted.
    """
    try:
        from backend.risk.scoring import score_detections
    except ImportError as exc:
        raise RuntimeError(
            "Could not import backend.risk.scoring.score_detections. "
            "Run the application from the project root."
        ) from exc

    scored_result = score_detections(detections)

    if isinstance(scored_result, dict):
        scored_items = scored_result.get(
            "detections",
            detections,
        )

        if not isinstance(scored_items, list):
            scored_items = detections

        scored_items = [
            item for item in scored_items
            if isinstance(item, dict)
        ]

        scores = [
            _clamp_score(item.get("risk_score", 0.0))
            for item in scored_items
        ]

        supplied_final = scored_result.get(
            "final_risk_score",
            scored_result.get("risk_score", None),
        )

        if supplied_final is None:
            final_score = max(scores, default=0.0)
        else:
            final_score = _clamp_score(supplied_final)

        supplied_average = scored_result.get(
            "average_risk_score",
            None,
        )

        if supplied_average is None:
            average_score = (
                sum(scores) / len(scores)
                if scores
                else 0.0
            )
        else:
            average_score = _clamp_score(supplied_average)

    elif isinstance(scored_result, list):
        scored_items = [
            item for item in scored_result
            if isinstance(item, dict)
        ]

        scores = [
            _clamp_score(item.get("risk_score", 0.0))
            for item in scored_items
        ]

        average_score = (
            sum(scores) / len(scores)
            if scores
            else 0.0
        )

        # The project uses the highest individual detection score
        # as the final file risk score.
        final_score = max(scores, default=0.0)

    else:
        raise TypeError(
            "score_detections() must return a dictionary or a list."
        )

    final_score = _clamp_score(final_score)
    average_score = _clamp_score(average_score)

    return {
        "detections": scored_items,
        "total_detections": len(scored_items),
        "average_risk_score": round(average_score, 2),
        "final_risk_score": round(final_score, 2),
        "risk_score": round(final_score, 2),
        "risk_level": _risk_level(final_score),
    }


# ============================================================
# FILE ANALYSIS
# ============================================================

def analyze_supported_file(
    file_path: str | Path,
    analyze_text_function=None,
    analyze_image_function=None,
    include_embedded_images: bool = True,
    render_pdf: bool = True,
    pdf_dpi: int = 150,
) -> Dict[str, Any]:
    """
    Analyze a supported document using PriviScan.

    Parameters
    ----------
    file_path:
        Path to PDF, DOCX, PPTX, XLSX, TXT, CSV, JSON, or MD.

    analyze_text_function:
        Existing PriviScan analyze_text function.

    analyze_image_function:
        Existing PriviScan analyze_image function.

    include_embedded_images:
        Analyze embedded images inside DOCX/PPTX documents.

    render_pdf:
        Render PDF pages for OCR when PyMuPDF is available.

    pdf_dpi:
        Rendering resolution used for scanned PDF pages.

    Returns
    -------
    dict
        Structured document-analysis result.
    """
    path = Path(file_path)

    extracted = extract_file_text(path)
    text = extracted.get("text", "")

    all_detections: List[Dict[str, Any]] = []
    text_result = None
    image_results = []
    rendered_pdf_results = []

    # --------------------------------------------------------
    # TEXT ANALYSIS
    # --------------------------------------------------------

    if text.strip() and analyze_text_function is not None:
        try:
            text_result = analyze_text_function(text)

            for detection in _get_detections(text_result):
                all_detections.append(
                    normalize_detection(
                        detection,
                        source="document_text",
                    )
                )

        except Exception as exc:
            print(
                f"WARNING: Text analysis failed for "
                f"{path.name}: {exc}"
            )

    # --------------------------------------------------------
    # EMBEDDED DOCX/PPTX IMAGES
    # --------------------------------------------------------

    if (
        include_embedded_images
        and analyze_image_function is not None
        and get_extension(path) in {".docx", ".pptx"}
    ):
        embedded_images = extract_embedded_images(path)

        for image_index, item in enumerate(
            embedded_images,
            start=1,
        ):
            temp_image_path = None

            try:
                with tempfile.NamedTemporaryFile(
                    suffix=Path(item["name"]).suffix,
                    delete=False,
                ) as temp_file:
                    temp_file.write(item["bytes"])
                    temp_image_path = Path(temp_file.name)

                image_result = analyze_image_function(
                    str(temp_image_path)
                )

                image_results.append({
                    "source": item["archive_path"],
                    "result": image_result,
                })

                for detection in _get_detections(image_result):
                    all_detections.append(
                        normalize_detection(
                            detection,
                            source="embedded_image",
                            source_location=item["archive_path"],
                        )
                    )

            except Exception as exc:
                print(
                    f"WARNING: Could not analyze embedded "
                    f"image {image_index}: {exc}"
                )

            finally:
                if temp_image_path is not None:
                    try:
                        temp_image_path.unlink(missing_ok=True)
                    except OSError:
                        pass

    # --------------------------------------------------------
    # SCANNED PDF PAGES
    # --------------------------------------------------------

    if (
        get_extension(path) == ".pdf"
        and render_pdf
        and analyze_image_function is not None
        and fitz is not None
        and len(text.strip()) < 50
    ):
        try:
            with tempfile.TemporaryDirectory(
                prefix="priviscan_pdf_"
            ) as temp_directory:
                rendered_pages = render_pdf_pages(
                    path,
                    temp_directory,
                    dpi=pdf_dpi,
                )

                for page_number, page_path in enumerate(
                    rendered_pages,
                    start=1,
                ):
                    try:
                        page_result = analyze_image_function(
                            str(page_path)
                        )

                        rendered_pdf_results.append({
                            "page": page_number,
                            "result": page_result,
                        })

                        for detection in _get_detections(
                            page_result
                        ):
                            all_detections.append(
                                normalize_detection(
                                    detection,
                                    source="pdf_page_image",
                                    source_location=f"page {page_number}",
                                )
                            )

                    except Exception as exc:
                        print(
                            f"WARNING: Could not analyze PDF "
                            f"page {page_number}: {exc}"
                        )

        except Exception as exc:
            print(
                f"WARNING: Could not render PDF {path.name}: {exc}"
            )

    # --------------------------------------------------------
    # MERGE DETECTIONS
    # --------------------------------------------------------

    all_detections = merge_detections(all_detections)

    # --------------------------------------------------------
    # RISK SCORING
    # --------------------------------------------------------

    scoring_error = None

    try:
        scoring_result = _score_detections_compat(
            all_detections
        )

        all_detections = scoring_result["detections"]
        average_risk_score = scoring_result[
            "average_risk_score"
        ]
        final_risk_score = scoring_result[
            "final_risk_score"
        ]
        risk_level = scoring_result["risk_level"]

    except Exception as exc:
        # Preserve analysis results even if the scorer is unavailable.
        # This fallback does not invent risk scores.
        scoring_error = str(exc)

        existing_scores = [
            _clamp_score(item.get("risk_score", 0.0))
            for item in all_detections
        ]

        average_risk_score = (
            sum(existing_scores) / len(existing_scores)
            if existing_scores
            else 0.0
        )

        final_risk_score = max(existing_scores, default=0.0)
        risk_level = _risk_level(final_risk_score)

        print(
            f"WARNING: Risk scoring could not be applied "
            f"to {path.name}: {exc}"
        )

    # Always derive the displayed level from the final score.
    final_risk_score = _clamp_score(final_risk_score)
    average_risk_score = _clamp_score(average_risk_score)
    risk_level = _risk_level(final_risk_score)

    # --------------------------------------------------------
    # RESULT
    # --------------------------------------------------------

    result = {
        "file_path": str(path),
        "file_name": path.name,
        "file_type": extracted.get("file_type"),
        "extension": extracted.get("extension"),
        "mime_type": extracted.get("mime_type"),
        "text": text,
        "text_length": len(text),
        "metadata": {
            key: value
            for key, value in extracted.items()
            if key not in {
                "text",
                "file_path",
                "file_name",
                "file_type",
                "extension",
                "mime_type",
            }
        },
        "text_analysis": text_result,
        "embedded_image_analysis": image_results,
        "pdf_page_analysis": rendered_pdf_results,
        "detections": all_detections,
        "total_detections": len(all_detections),
        "average_risk_score": round(average_risk_score, 2),
        "final_risk_score": round(final_risk_score, 2),
        "risk_score": round(final_risk_score, 2),
        "risk_level": risk_level,
    }

    if scoring_error is not None:
        result["scoring_error"] = scoring_error

    return result


# ============================================================
# SIMPLE TEXT-ONLY ANALYSIS
# ============================================================

def analyze_file_text_only(
    file_path: str | Path,
    analyze_text_function,
) -> Dict[str, Any]:
    """
    Extract document text and pass it to PriviScan's text pipeline.
    """
    extracted = extract_file_text(file_path)
    text = extracted.get("text", "")

    if text.strip():
        result = analyze_text_function(text)
    else:
        result = None

    detections = _get_detections(result)

    return {
        **extracted,
        "text_analysis": result,
        "detections": detections,
        "total_detections": len(detections),
    }


# ============================================================
# CLI TEST
# ============================================================

def main() -> None:
    """Run a command-line text extraction test."""
    import argparse

    parser = argparse.ArgumentParser(
        description="PriviScan supported-file text extraction test"
    )

    parser.add_argument(
        "file",
        help="Path to PDF/DOCX/PPTX/XLSX/TXT/CSV/JSON/MD",
    )

    args = parser.parse_args()
    path = Path(args.file)

    if not path.exists():
        print(f"ERROR: File not found: {path}")
        raise SystemExit(1)

    try:
        result = extract_file_text(path)

    except Exception as exc:
        print(f"ERROR: {exc}")
        raise SystemExit(1)

    print()
    print("=" * 70)
    print("PRIVISCAN FILE SUPPORT TEST")
    print("=" * 70)
    print(f"File       : {result['file_name']}")
    print(f"Type       : {result['file_type']}")
    print(f"Extension  : {result['extension']}")
    print(f"MIME       : {result['mime_type']}")
    print(f"Text chars : {len(result.get('text', ''))}")

    print()
    print("-" * 70)
    print("EXTRACTED TEXT")
    print("-" * 70)

    text = result.get("text", "")

    if text:
        print(text)
    else:
        print("[No selectable text extracted]")

    print()
    print("=" * 70)


if __name__ == "__main__":
    main()

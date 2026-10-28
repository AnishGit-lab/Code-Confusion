"""
Create PriviScan document-format test files.

Creates:

    PDF
    DOCX
    PPTX
    XLSX
    TXT
    CSV
    JSON
    MD

All files contain synthetic test information only.
"""

from pathlib import Path
import csv
import json


OUTPUT_DIR = (
    Path(__file__).resolve().parent.parent
    / "dataset"
    / "file_tests"
)


TEST_DATA = {
    "name": "Anish Kandel",
    "email": "anish.kandel@example.com",
    "phone": "+82 10-1234-5678",
    "address": "123 Example Street, Seoul",
    "student_id": "2024891090",
    "employee_id": "EMP202489",
    "government_id": "P123456789",
    "account_number": "123456789012",
    "password": "MySecretPassword123",
    "api_key": "sk-test-123456789",
}


def create_txt():
    path = OUTPUT_DIR / "priviscan_test.txt"

    text = f"""
PriviScan File Test

Name: {TEST_DATA["name"]}
Email: {TEST_DATA["email"]}
Phone: {TEST_DATA["phone"]}
Address: {TEST_DATA["address"]}
Student ID: {TEST_DATA["student_id"]}
Employee ID: {TEST_DATA["employee_id"]}
Government ID: {TEST_DATA["government_id"]}
Account Number: {TEST_DATA["account_number"]}
Password: {TEST_DATA["password"]}
API Key: {TEST_DATA["api_key"]}
""".strip()

    path.write_text(
        text,
        encoding="utf-8",
    )

    return path


def create_md():
    path = OUTPUT_DIR / "priviscan_test.md"

    text = f"""# PriviScan File Test

## Personal Information

- Name: {TEST_DATA["name"]}
- Email: {TEST_DATA["email"]}
- Phone: {TEST_DATA["phone"]}
- Address: {TEST_DATA["address"]}
- Student ID: {TEST_DATA["student_id"]}

## Sensitive Credentials

- Government ID: {TEST_DATA["government_id"]}
- Account Number: {TEST_DATA["account_number"]}
- Password: {TEST_DATA["password"]}
- API Key: {TEST_DATA["api_key"]}
"""

    path.write_text(
        text,
        encoding="utf-8",
    )

    return path


def create_csv():
    path = OUTPUT_DIR / "priviscan_test.csv"

    rows = [
        [
            "Name",
            "Email",
            "Phone",
            "Student ID",
            "Account Number",
        ],
        [
            TEST_DATA["name"],
            TEST_DATA["email"],
            TEST_DATA["phone"],
            TEST_DATA["student_id"],
            TEST_DATA["account_number"],
        ],
    ]

    with path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as file:
        writer = csv.writer(file)
        writer.writerows(rows)

    return path


def create_json():
    path = OUTPUT_DIR / "priviscan_test.json"

    data = {
        "name": TEST_DATA["name"],
        "email": TEST_DATA["email"],
        "phone": TEST_DATA["phone"],
        "address": TEST_DATA["address"],
        "student_id": TEST_DATA["student_id"],
        "government_id": TEST_DATA["government_id"],
        "account_number": TEST_DATA["account_number"],
        "password": TEST_DATA["password"],
        "api_key": TEST_DATA["api_key"],
    }

    path.write_text(
        json.dumps(
            data,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    return path


def create_docx():
    from docx import Document

    path = OUTPUT_DIR / "priviscan_test.docx"

    document = Document()

    document.add_heading(
        "PriviScan File Test",
        level=1,
    )

    values = [
        ("Name", TEST_DATA["name"]),
        ("Email", TEST_DATA["email"]),
        ("Phone", TEST_DATA["phone"]),
        ("Address", TEST_DATA["address"]),
        ("Student ID", TEST_DATA["student_id"]),
        ("Employee ID", TEST_DATA["employee_id"]),
        ("Government ID", TEST_DATA["government_id"]),
        ("Account Number", TEST_DATA["account_number"]),
        ("Password", TEST_DATA["password"]),
        ("API Key", TEST_DATA["api_key"]),
    ]

    for label, value in values:
        document.add_paragraph(
            f"{label}: {value}"
        )

    table = document.add_table(
        rows=1,
        cols=2,
    )

    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Value"

    for label, value in values[:5]:
        row = table.add_row()

        row.cells[0].text = label
        row.cells[1].text = value

    document.save(
        str(path)
    )

    return path


def create_pptx():
    from pptx import Presentation

    path = OUTPUT_DIR / "priviscan_test.pptx"

    presentation = Presentation()

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide.shapes.title.text = (
        "PriviScan File Test"
    )

    body = slide.placeholders[1]

    body.text = (
        f"Name: {TEST_DATA['name']}\n"
        f"Email: {TEST_DATA['email']}\n"
        f"Phone: {TEST_DATA['phone']}\n"
        f"Address: {TEST_DATA['address']}\n"
        f"Student ID: {TEST_DATA['student_id']}\n"
        f"Employee ID: {TEST_DATA['employee_id']}\n"
        f"Government ID: {TEST_DATA['government_id']}\n"
        f"Account Number: {TEST_DATA['account_number']}\n"
        f"Password: {TEST_DATA['password']}\n"
        f"API Key: {TEST_DATA['api_key']}"
    )

    # Second slide
    slide2 = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide2.shapes.title.text = (
        "Additional Sensitive Data"
    )

    body2 = slide2.placeholders[1]

    body2.text = (
        f"Email: {TEST_DATA['email']}\n"
        f"Phone: {TEST_DATA['phone']}\n"
        f"Student ID: {TEST_DATA['student_id']}"
    )

    presentation.save(
        str(path)
    )

    return path


def create_xlsx():
    from openpyxl import Workbook

    path = OUTPUT_DIR / "priviscan_test.xlsx"

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Privacy Test"

    worksheet.append(
        [
            "Field",
            "Value",
        ]
    )

    for label, value in [
        ("Name", TEST_DATA["name"]),
        ("Email", TEST_DATA["email"]),
        ("Phone", TEST_DATA["phone"]),
        ("Address", TEST_DATA["address"]),
        ("Student ID", TEST_DATA["student_id"]),
        ("Employee ID", TEST_DATA["employee_id"]),
        ("Government ID", TEST_DATA["government_id"]),
        ("Account Number", TEST_DATA["account_number"]),
        ("Password", TEST_DATA["password"]),
        ("API Key", TEST_DATA["api_key"]),
    ]:
        worksheet.append(
            [
                label,
                value,
            ]
        )

    workbook.save(
        str(path)
    )

    return path


def create_pdf():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = OUTPUT_DIR / "priviscan_test.pdf"

    pdf = canvas.Canvas(
        str(path),
        pagesize=letter,
    )

    width, height = letter

    y = height - 60

    pdf.setFont(
        "Helvetica-Bold",
        16,
    )

    pdf.drawString(
        50,
        y,
        "PriviScan File Test",
    )

    y -= 40

    pdf.setFont(
        "Helvetica",
        11,
    )

    values = [
        ("Name", TEST_DATA["name"]),
        ("Email", TEST_DATA["email"]),
        ("Phone", TEST_DATA["phone"]),
        ("Address", TEST_DATA["address"]),
        ("Student ID", TEST_DATA["student_id"]),
        ("Employee ID", TEST_DATA["employee_id"]),
        ("Government ID", TEST_DATA["government_id"]),
        ("Account Number", TEST_DATA["account_number"]),
        ("Password", TEST_DATA["password"]),
        ("API Key", TEST_DATA["api_key"]),
    ]

    for label, value in values:
        pdf.drawString(
            50,
            y,
            f"{label}: {value}",
        )

        y -= 30

        if y < 60:
            pdf.showPage()

            pdf.setFont(
                "Helvetica",
                11,
            )

            y = height - 60

    pdf.save()

    return path


def main():
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    creators = [
        create_txt,
        create_md,
        create_csv,
        create_json,
        create_docx,
        create_pptx,
        create_xlsx,
        create_pdf,
    ]

    print()
    print("=" * 70)
    print("PRIVISCAN FILE TEST DATASET")
    print("=" * 70)

    for creator in creators:
        try:
            path = creator()

            print(
                f"Created: {path}"
            )

        except Exception as exc:
            print(
                f"FAILED: "
                f"{creator.__name__}: "
                f"{exc}"
            )

    print()
    print(
        f"Location: {OUTPUT_DIR}"
    )

    print()
    print(
        "File dataset generation complete."
    )


if __name__ == "__main__":
    main()